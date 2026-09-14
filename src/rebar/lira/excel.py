"""Строгий read-only импорт трёх XLSX профиля А101, без исполнения формул.

Инструкция А101: первая строка КЭ — полная потребность. Никакого усреднения,
автоподбора фона, толщины Revit или назначения фаз здесь нет.
"""
from __future__ import annotations

from contextlib import contextmanager
import hashlib
from io import BytesIO
import math
from pathlib import Path
import re
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.xml import DEFUSEDXML
from shapely.geometry import Polygon
from shapely.strtree import STRtree

from .models import LiraElement, LiraGroup, LiraNode, LiraPlate, LiraSource

MAX_FILE_BYTES = 32 * 1024**2
MAX_EXPANDED_BYTES = 128 * 1024**2
MAX_ROWS = 100_000
MAX_ELEMENTS = 10_000


def _text(value) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _number(value, where: str, *, positive: bool = False) -> float:
    if (isinstance(value, bool) or not isinstance(value, (int, float))
            or not math.isfinite(value) or (positive and value <= 0)):
        raise ValueError(f"{where}: требуется {'положительное ' if positive else ''}конечное число")
    return float(value)


def _integer(value, where: str) -> int:
    number = _number(value, where, positive=True)
    if not number.is_integer() or number > 2**53 - 1:
        raise ValueError(f"{where}: требуется точный положительный целый ID")
    return int(number)


@contextmanager
def _book(path: Path, role: str):
    if not DEFUSEDXML:
        raise ValueError("Для безопасного XLSX import требуется включённый defusedxml")
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"{role}: ожидается .xlsx")
    # Hash and parse the same bounded snapshot, not two reads of a changing file.
    with path.open("rb") as stream:
        data = stream.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise ValueError(f"{role}: XLSX превышает лимит 32 МиБ")
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if (len(entries) > 512 or sum(i.file_size for i in entries) > MAX_EXPANDED_BYTES
                    or len({i.filename for i in entries}) != len(entries)):
                raise ValueError(f"{role}: недопустимый размер/повторяющиеся записи XLSX")
            if any("vbaproject" in i.filename.lower() or "externallinks/" in i.filename.lower()
                   for i in entries):
                raise ValueError(f"{role}: макросы/внешние связи не поддержаны")
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, ParseError, DefusedXmlException, InvalidFileException) as error:
        raise ValueError(f"{role}: повреждённый XLSX") from error
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError(f"{role}: требуется один явно выделенный лист")
        sheet = workbook.worksheets[0]
        if sheet.sheet_state != "visible":
            raise ValueError(f"{role}: скрытый лист не является явным источником")
        # LIRA omits dimensions; other exporters may write stale A1:A1 dimensions.
        sheet.reset_dimensions()
        rows = []
        for number, cells in enumerate(sheet.iter_rows(), 1):
            if number > MAX_ROWS or len(cells) > 64:
                raise ValueError(f"{role}: превышен лимит строк/столбцов")
            if any(c.data_type in ("f", "e") for c in cells):
                raise ValueError(f"{role}, строка {number}: формула/ошибка вместо исходных данных")
            rows.append(tuple(c.value for c in cells))
        yield rows, LiraSource(role, path.name, hashlib.sha256(data).hexdigest(), sheet.title, len(rows))
    except (ParseError, DefusedXmlException) as error:
        raise ValueError(f"{role}: повреждённый или небезопасный XML листа") from error
    finally:
        workbook.close()


def _header(rows, index: int, expected: tuple[str, ...], role: str):
    if len(rows) <= index or tuple(_text(v) for v in rows[index][:len(expected)]) != expected:
        raise ValueError(f"{role}: не распознан заголовок строки {index + 1}")


def _nodes(rows):
    _header(rows, 0, ("Таблица узлов",), "Узлы")
    if len(rows) < 4 or len(rows[2]) < 4 or _text(rows[2][0]) != "№ узла":
        raise ValueError("Узлы: не распознан заголовок координат")
    units = []
    for axis, value in zip("XYZ", rows[2][1:4]):
        match = re.fullmatch(rf"{axis}\s*\((м|мм)\)", _text(value))
        if match is None:
            raise ValueError("Узлы: требуются явно указанные единицы X/Y/Z (м или мм)")
        units.append(match[1])
    if len(units) != 3 or len(set(units)) != 1:
        raise ValueError("Узлы: единицы X/Y/Z должны совпадать")
    unit = {"м": "m", "мм": "mm"}[units[0]]
    scale = 1000 if unit == "m" else 1
    result = {}
    for number, row in enumerate(rows[3:], 4):
        if not any(v is not None for v in row):
            continue
        if len(row) < 4:
            raise ValueError(f"Узлы, строка {number}: отсутствуют XYZ")
        identifier = _integer(row[0], f"Узлы!A{number}")
        if identifier in result:
            raise ValueError(f"Узлы: повтор ID {identifier}")
        coordinates = tuple(_number(v, f"Узлы!{axis}{number}") * scale
                            for axis, v in zip("BCD", row[1:4]))
        if any(not math.isfinite(v) or abs(v) > 1e9 for v in coordinates):
            raise ValueError(f"Узел {identifier}: координаты вне поддержанного диапазона ±1e9 мм")
        result[identifier] = LiraNode(identifier, coordinates, number)
    if not result:
        raise ValueError("Узлы: нет данных")
    return result, unit


def _elements(rows, nodes):
    _header(rows, 0, ("Таблица элементов",), "Элементы")
    _header(rows, 2, ("№ элем", "Тип элем", "Кол.сечений", "Тип жестк", "Угол м.осей",
                      "AX н", "AX к", "№№ узлов"), "Элементы")
    result = {}
    for number, row in enumerate(rows[3:], 4):
        if not any(v is not None for v in row):
            continue
        if len(row) < 8:
            raise ValueError(f"Элементы, строка {number}: неполная запись")
        identifier = _integer(row[0], f"Элементы!A{number}")
        if identifier in result:
            raise ValueError(f"Элементы: повтор ID {identifier}")
        kind = _integer(row[1], f"Элементы!B{number}")
        if kind not in (42, 44) or row[2] != 1 or isinstance(row[2], bool):
            raise ValueError(f"КЭ {identifier}: поддержаны оболочки 42/44 с одним сечением")
        stiffness = _integer(row[3], f"Элементы!D{number}")
        angle = None if row[4] == "-" else _number(row[4], f"Элементы!E{number}")
        if any(value not in (None, "-", 0) for value in row[5:7]):
            raise ValueError(f"КЭ {identifier}: жёсткие вставки не поддержаны")
        if not isinstance(row[7], str) or not re.fullmatch(r"\s*\d+(?:\s*,\s*\d+)+\s*", row[7]):
            raise ValueError(f"КЭ {identifier}: неверный список узлов")
        ids = tuple(int(v.strip()) for v in row[7].split(","))
        if len(ids) != {42: 3, 44: 4}[kind] or len(set(ids)) != len(ids):
            raise ValueError(f"КЭ {identifier}: число/уникальность узлов не соответствует типу")
        missing = set(ids) - nodes.keys()
        if missing:
            raise ValueError(f"КЭ {identifier}: отсутствуют узлы {sorted(missing)}")
        vertices = tuple(nodes[n].xyz_mm for n in ids)
        if len({point[:2] for point in vertices}) != len(ids):
            raise ValueError(f"КЭ {identifier}: разные ID с совпадающими вершинами XY")
        # Preserve original IDs/vertices, use the profile's explicit FE perimeter order.
        order = (0, 1, 3, 2) if kind == 44 else (0, 1, 2)
        poly = Polygon([vertices[i][:2] for i in order])
        if (not poly.is_valid or poly.area <= 1e-6
                or poly.convex_hull.area - poly.area > max(1e-6, poly.area * 1e-10)):
            raise ValueError(f"КЭ {identifier}: вырожденная/невыпуклая геометрия XY")
        result[identifier] = (kind, stiffness, angle, ids, vertices, number)
    if not result or len(result) > MAX_ELEMENTS:
        raise ValueError("Элементы: требуется от 1 до 10000 КЭ")
    return result


def _group(text: str, row: int) -> LiraGroup:
    number = r"\d+(?:[.,]\d+)?"
    match = re.fullmatch(
        rf"(\d+)\s*-\s*Оболочка\s*/\s*h\s*=\s*({number})\s*см\s*/\s*"
        r"Бетон\s+([^/]+)\s*/\s*Арматура:\s*продольная\s+Ax:\s*([^,]+),\s*Ay:\s*([^/]+)"
        rf"/\s*поперечная\s+[^/]+/\s*Шаг арматурных стержней\s+({number})\s*мм",
        _text(text),
    )
    if match is None:
        raise ValueError(f"Арматура, строка {row}: не распознан заголовок группы оболочек")
    thickness = _number(float(match[2].replace(",", ".")), f"Группа, строка {row}: h", positive=True) * 10
    step = _number(float(match[6].replace(",", ".")), f"Группа, строка {row}: шаг", positive=True)
    return LiraGroup(_integer(int(match[1]), f"Группа, строка {row}"), thickness,
                     match[3].strip(), match[4].strip(), match[5].strip(), step, row, text)


def _reinforcement(rows):
    _header(rows, 0, ("ГР", "Элемент", "AS1", "AS2", "AS3", "AS4", "ASW1", "ASW2",
                      "Кратк.", "Длит."), "Арматура")
    groups, result = {}, {}
    group = pending = None
    for number, row in enumerate(rows[1:], 2):
        if not any(v is not None for v in row):
            continue
        if isinstance(row[0], str) and all(v is None for v in row[1:]):
            if pending is not None:
                raise ValueError(f"КЭ {pending[0]}: отсутствует вторая строка перед новой группой")
            group = _group(row[0], number)
            if group.id in groups:
                raise ValueError(f"Арматура: повтор группы {group.id}")
            groups[group.id] = group
            continue
        if group is None or len(row) < 6:
            raise ValueError(f"Арматура, строка {number}: нет группы или AS1–AS4")
        if _integer(row[0], f"Арматура!A{number}") != group.id:
            raise ValueError(f"Арматура, строка {number}: ID группы отличается от заголовка")
        identifier = _integer(row[1], f"Арматура!B{number}")
        values = tuple(_number(v, f"Арматура!{c}{number}") for c, v in zip("CDEF", row[2:6]))
        if any(v < 0 for v in values):
            raise ValueError(f"КЭ {identifier}: отрицательная площадь армирования")
        if pending is None:
            if identifier in result:
                raise ValueError(f"Арматура: повтор пары КЭ {identifier}")
            pending = (identifier, values, number)
        else:
            if identifier != pending[0]:
                raise ValueError(f"КЭ {pending[0]}: вторая строка содержит другой ID {identifier}")
            if any(second > first for first, second in zip(pending[1], values)):
                raise ValueError(f"КЭ {identifier}: первая строка меньше второй; порядок/профиль не подтверждён")
            result[identifier] = (group.id, pending[1], values, (pending[2], number))
            pending = None
    if pending is not None:
        raise ValueError(f"КЭ {pending[0]}: отсутствует вторая строка")
    if not result:
        raise ValueError("Арматура: нет данных")
    used = {value[0] for value in result.values()}
    if used != groups.keys():
        raise ValueError("Арматура: группа не содержит КЭ")
    return groups, result


def _validate_mesh(elements):
    polygons = [Polygon(element.polygon_xy_mm) for element in elements]
    index = STRtree(polygons)
    for i, poly in enumerate(polygons):
        for j in index.query(poly, predicate="intersects"):
            if j <= i:
                continue
            if poly.intersection(polygons[j]).area > max(1e-6, min(poly.area, polygons[j].area) * 1e-10):
                raise ValueError(f"КЭ {elements[i].id}/{elements[j].id}: наложение геометрии")


def read_lira_excel(nodes_path: str | Path, elements_path: str | Path,
                    reinforcement_path: str | Path) -> LiraPlate:
    """Прочитать один горизонтальный комплект; любой потерянный КЭ — ошибка всего входа."""
    with _book(Path(nodes_path), "nodes") as (rows, node_source):
        nodes, unit = _nodes(rows)
    with _book(Path(elements_path), "elements") as (rows, element_source):
        geometries = _elements(rows, nodes)
    with _book(Path(reinforcement_path), "reinforcement") as (rows, reinforcement_source):
        groups, reinforcement = _reinforcement(rows)
    missing, extra = geometries.keys() - reinforcement.keys(), reinforcement.keys() - geometries.keys()
    if missing or extra:
        raise ValueError(f"Арматура/геометрия: отсутствуют КЭ {sorted(missing)[:20]}, лишние {sorted(extra)[:20]}")
    elements = tuple(LiraElement(
        identifier, *geometry[:5], *reinforcement[identifier][:3],
        geometry[5], reinforcement[identifier][3],
    ) for identifier, geometry in sorted(geometries.items()))
    used_ids = {n for e in elements for n in e.node_ids}
    used_nodes = [nodes[n] for n in sorted(used_ids)]
    bbox = tuple(tuple(fn(n.xyz_mm[a] for n in used_nodes) for a in range(3)) for fn in (min, max))
    if bbox[1][2] - bbox[0][2] > 1e-3:
        raise ValueError("Ожидается одна горизонтальная плита; разные отметки Z не сплющиваются")
    _validate_mesh(elements)
    return LiraPlate((node_source, element_source, reinforcement_source),
                     tuple(nodes[i] for i in sorted(nodes)), tuple(groups[i] for i in sorted(groups)),
                     elements, bbox, unit, tuple(sorted(nodes.keys() - used_ids)))
